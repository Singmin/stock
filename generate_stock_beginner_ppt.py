from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


OUT = Path(r"C:\Users\Singmin\Desktop\stock\炒股初学者入门课_案例图解版.pptx")

FONT = "Microsoft YaHei"
BG = RGBColor(18, 30, 43)
BG2 = RGBColor(26, 43, 60)
CARD = RGBColor(246, 249, 252)
TEXT = RGBColor(32, 45, 59)
MUTED = RGBColor(96, 111, 126)
WHITE = RGBColor(255, 255, 255)
RED = RGBColor(220, 68, 85)
GREEN = RGBColor(23, 166, 118)
GOLD = RGBColor(232, 177, 69)
BLUE = RGBColor(70, 142, 214)
CYAN = RGBColor(79, 190, 198)


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def rgb(hex_value):
    hex_value = hex_value.lstrip("#")
    return RGBColor(int(hex_value[0:2], 16), int(hex_value[2:4], 16), int(hex_value[4:6], 16))


def set_fill(shape, color, transparency=None):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if transparency is not None:
        shape.fill.transparency = transparency


def no_line(shape):
    shape.line.fill.background()


def add_textbox(slide, x, y, w, h, text, size=22, color=TEXT, bold=False,
                align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.08):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_bg(slide):
    rect = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    set_fill(rect, BG)
    no_line(rect)
    band = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.05))
    set_fill(band, BG2)
    no_line(band)
    stripe = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, Inches(7.18), prs.slide_width, Inches(0.32))
    set_fill(stripe, rgb("0F1B27"))
    no_line(stripe)


def add_header(slide, title, section=None, page=None):
    add_textbox(slide, 0.55, 0.20, 9.7, 0.52, title, 24, WHITE, True)
    if section:
        pill = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(10.35), Inches(0.23), Inches(2.25), Inches(0.42))
        set_fill(pill, rgb("27445E"))
        pill.line.color.rgb = rgb("385B78")
        add_textbox(slide, 10.42, 0.26, 2.1, 0.28, section, 10, rgb("BFD3E5"), False, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 0)
    if page is not None:
        add_textbox(slide, 12.55, 7.20, 0.5, 0.18, f"{page:02d}", 8, rgb("9FB2C3"), False, PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE, 0)


def add_footer(slide, note):
    add_textbox(slide, 0.55, 7.17, 11.9, 0.22, f"讲师提示：{note}", 7.5, rgb("B8C5D2"), False, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE, 0)


def add_card(slide, x, y, w, h, fill=CARD, line=rgb("D9E2EA")):
    card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    card.adjustments[0] = 0.08
    set_fill(card, fill)
    card.line.color.rgb = line
    card.line.width = Pt(0.7)
    return card


def add_bullets(slide, x, y, w, h, bullets, size=16, color=TEXT, gap=5, bullet=True):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.04)
    for idx, text in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = text
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(gap)
        if bullet:
            p.level = 0
            p._p.get_or_add_pPr().set("marL", "171450")
            p._p.get_or_add_pPr().set("indent", "-114300")
    return box


def base_slide(title, section, note, page):
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_header(slide, title, section, page)
    add_footer(slide, note)
    return slide


def cover_slide(page):
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_textbox(slide, 0.72, 0.62, 4.2, 0.4, "股票投资入门 · 教育课件", 15, rgb("BFD3E5"), True)
    add_textbox(slide, 0.68, 1.38, 8.9, 1.25, "炒股初学者入门课", 46, WHITE, True)
    add_textbox(slide, 0.72, 2.68, 8.8, 0.5, "从交易规则到K线、均线、风控与复盘", 22, rgb("D8E5EF"))
    add_card(slide, 0.75, 4.68, 5.7, 1.0, rgb("20364B"), rgb("385B78"))
    add_textbox(slide, 1.05, 4.90, 5.1, 0.42, "适合零基础、刚开户、想建立正确交易框架的新手", 17, WHITE, False)

    # Decorative market board
    board = add_card(slide, 8.15, 1.05, 4.35, 4.95, rgb("F7FAFD"), rgb("DBE4EC"))
    add_textbox(slide, 8.47, 1.30, 3.6, 0.3, "Market Basics", 18, TEXT, True)
    x0, y0 = 8.55, 2.0
    data = [
        ("A股", "+1.2%", RED, [1.1, 1.5, 1.3, 1.9, 2.2, 2.6, 2.4]),
        ("港股", "-0.4%", GREEN, [2.5, 2.3, 2.2, 2.4, 2.0, 1.8, 1.9]),
        ("美股", "+0.8%", RED, [1.6, 1.7, 2.1, 2.0, 2.3, 2.5, 2.8]),
    ]
    for r, (name, pct, col, vals) in enumerate(data):
        y = y0 + r * 1.05
        add_textbox(slide, x0, y, 0.65, 0.25, name, 12, TEXT, True, margin=0)
        add_textbox(slide, x0 + 0.7, y, 0.7, 0.25, pct, 12, col, True, PP_ALIGN.RIGHT, margin=0)
        for i in range(len(vals) - 1):
            x1 = Inches(x0 + 1.65 + i * 0.25)
            y1 = Inches(y + 0.52 - vals[i] * 0.12)
            x2 = Inches(x0 + 1.65 + (i + 1) * 0.25)
            y2 = Inches(y + 0.52 - vals[i + 1] * 0.12)
            line = slide.shapes.add_connector(1, x1, y1, x2, y2)
            line.line.color.rgb = col
            line.line.width = Pt(2.0)
    add_textbox(slide, 0.55, 7.17, 11.9, 0.22, "资料口径截至 2026-06-12；本课件仅用于投资者教育，不构成投资建议。", 7.5, rgb("B8C5D2"), False)
    return slide


def section_slide(title, subtitle, items, page):
    slide = prs.slides.add_slide(blank)
    add_bg(slide)
    add_textbox(slide, 0.78, 1.08, 8.5, 0.65, title, 34, WHITE, True)
    add_textbox(slide, 0.82, 1.82, 9.6, 0.42, subtitle, 18, rgb("D4E2EF"))
    for idx, item in enumerate(items):
        y = 3.0 + idx * 0.72
        dot = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(0.92), Inches(y + 0.08), Inches(0.16), Inches(0.16))
        set_fill(dot, [RED, GREEN, GOLD, CYAN][idx % 4])
        no_line(dot)
        add_textbox(slide, 1.18, y, 8.9, 0.35, item, 18, WHITE, False)
    add_textbox(slide, 12.55, 7.20, 0.5, 0.18, f"{page:02d}", 8, rgb("9FB2C3"), False, PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE, 0)
    return slide


def text_slide(title, section, bullets, note, page, side_title=None, side_items=None):
    slide = base_slide(title, section, note, page)
    add_card(slide, 0.72, 1.28, 7.25, 5.45)
    add_bullets(slide, 1.05, 1.65, 6.55, 4.75, bullets[:3], 20, gap=9)
    add_card(slide, 8.35, 1.28, 4.25, 5.45, rgb("EDF4F9"))
    if side_title:
        add_textbox(slide, 8.68, 1.58, 3.55, 0.32, side_title, 18, TEXT, True)
    if side_items:
        y = 2.15
        for idx, item in enumerate(side_items):
            color = [BLUE, RED, GREEN, GOLD, CYAN][idx % 5]
            badge = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(8.72), Inches(y), Inches(0.62), Inches(0.36))
            set_fill(badge, color)
            no_line(badge)
            add_textbox(slide, 8.72, y + 0.05, 0.62, 0.18, f"{idx + 1}", 9, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 0)
            add_textbox(slide, 9.55, y - 0.01, 2.75, 0.42, item, 13.5, TEXT)
            y += 0.82
    return slide


def table_slide(title, section, headers, rows, note, page, foot=None):
    slide = base_slide(title, section, note, page)
    rows_count = len(rows) + 1
    cols_count = len(headers)
    table_shape = slide.shapes.add_table(rows_count, cols_count, Inches(0.65), Inches(1.32), Inches(12.0), Inches(5.45))
    table = table_shape.table
    for c, header in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = header
        set_fill(cell, rgb("2B4B66"))
        for p in cell.text_frame.paragraphs:
            p.font.name = FONT
            p.font.bold = True
            p.font.size = Pt(12)
            p.font.color.rgb = WHITE
            p.alignment = PP_ALIGN.CENTER
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = val
            set_fill(cell, CARD if r % 2 else rgb("EEF4F8"))
            cell.margin_left = Inches(0.05)
            cell.margin_right = Inches(0.05)
            for p in cell.text_frame.paragraphs:
                p.font.name = FONT
                p.font.size = Pt(9.8 if len(val) < 28 else 8.5)
                p.font.color.rgb = TEXT
                p.alignment = PP_ALIGN.CENTER if c == 0 else PP_ALIGN.LEFT
    for c, width in enumerate([1.3, 2.15, 1.7, 2.15, 2.0, 2.7][:cols_count]):
        table.columns[c].width = Inches(width)
    if foot:
        add_textbox(slide, 0.72, 6.72, 11.8, 0.25, foot, 8.5, rgb("C5D2DF"))
    return slide


def draw_candle(slide, x, y, open_p, close_p, high_p, low_p, scale=0.045, label=None):
    col = RED if close_p >= open_p else GREEN
    center = Inches(x)
    top = Inches(y + (100 - high_p) * scale)
    bottom = Inches(y + (100 - low_p) * scale)
    wick = slide.shapes.add_connector(1, center, top, center, bottom)
    wick.line.color.rgb = col
    wick.line.width = Pt(1.5)
    body_top_p = max(open_p, close_p)
    body_bottom_p = min(open_p, close_p)
    body_y = y + (100 - body_top_p) * scale
    body_h = max(0.12, (body_top_p - body_bottom_p) * scale)
    body = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x - 0.09), Inches(body_y), Inches(0.18), Inches(body_h))
    set_fill(body, col if close_p != open_p else WHITE)
    body.line.color.rgb = col
    body.line.width = Pt(1)
    if label:
        add_textbox(slide, x - 0.34, y + 3.30, 0.68, 0.28, label, 8.5, MUTED, False, PP_ALIGN.CENTER, margin=0)


def kline_slide(title, bullets, note, page):
    slide = base_slide(title, "K线图入门", note, page)
    add_card(slide, 0.70, 1.28, 5.15, 5.45)
    add_bullets(slide, 1.02, 1.60, 4.55, 4.9, bullets[:3], 18, gap=8)
    add_card(slide, 6.25, 1.28, 6.35, 5.45, rgb("F7FAFC"))
    add_textbox(slide, 6.55, 1.55, 5.7, 0.32, "K线示意：看价格区间，也看所在位置", 17, TEXT, True)
    # Axis
    axis = slide.shapes.add_connector(1, Inches(6.75), Inches(5.65), Inches(12.0), Inches(5.65))
    axis.line.color.rgb = rgb("B7C7D6")
    vals = [
        (74, 88, 91, 70, "阳线"),
        (87, 76, 92, 73, "阴线"),
        (80, 80, 93, 68, "十字"),
        (70, 91, 94, 68, "大阳"),
        (92, 69, 95, 66, "大阴"),
        (78, 86, 98, 76, "上影"),
        (84, 77, 86, 61, "下影"),
    ]
    for i, item in enumerate(vals):
        draw_candle(slide, 7.15 + i * 0.72, 2.10, item[0], item[1], item[2], item[3], label=item[4])
    return slide


def ma_chart_slide(page):
    slide = base_slide("均线如何帮助判断趋势", "均线系统", "均线描述的是平均成本与趋势，不是自动买卖按钮。", page)
    add_card(slide, 0.72, 1.25, 12.0, 5.45, rgb("F7FAFC"))
    add_textbox(slide, 1.05, 1.50, 7.2, 0.32, "示意：价格上行，短期均线逐步站上中长期均线", 17, TEXT, True)
    # chart area
    x0, y0, w, h = 1.05, 2.05, 7.3, 3.85
    for i in range(5):
        y = y0 + i * h / 4
        line = slide.shapes.add_connector(1, Inches(x0), Inches(y), Inches(x0 + w), Inches(y))
        line.line.color.rgb = rgb("E0E7EE")
        line.line.width = Pt(0.7)
    price = [2.8, 2.7, 3.1, 3.0, 3.5, 3.8, 3.6, 4.3, 4.7, 4.5, 5.2, 5.6, 5.3, 6.0]
    ma5 = [2.75, 2.82, 2.9, 3.1, 3.22, 3.45, 3.55, 3.8, 4.05, 4.25, 4.55, 4.85, 5.08, 5.45]
    ma20 = [3.0, 3.02, 3.06, 3.12, 3.2, 3.32, 3.45, 3.6, 3.78, 3.95, 4.15, 4.4, 4.62, 4.9]
    def plot(vals, col, width=2.2):
        mn, mx = 2.3, 6.2
        pts = []
        for i, v in enumerate(vals):
            x = x0 + i * w / (len(vals) - 1)
            y = y0 + h - (v - mn) / (mx - mn) * h
            pts.append((x, y))
        for i in range(len(pts) - 1):
            l = slide.shapes.add_connector(1, Inches(pts[i][0]), Inches(pts[i][1]), Inches(pts[i + 1][0]), Inches(pts[i + 1][1]))
            l.line.color.rgb = col
            l.line.width = Pt(width)
    plot(price, rgb("22364A"), 2.5)
    plot(ma5, RED, 2)
    plot(ma20, BLUE, 2)
    add_textbox(slide, 1.1, 6.02, 6.8, 0.25, "黑线=价格  红线=短期均线  蓝线=中期均线", 10, MUTED)
    add_card(slide, 8.75, 2.0, 3.45, 3.85, rgb("EEF4F8"))
    add_bullets(slide, 9.02, 2.34, 2.9, 3.1, [
        "短期均线：反应快，噪声也多",
        "中期均线：适合观察一段趋势",
        "长期均线：更像市场大方向",
        "跌破均线不等于立刻卖，需结合成交量和计划",
    ], 13.5)
    return slide


def quadrant_slide(page):
    slide = base_slide("量价关系：先看参与度，再看方向", "成交量与指标", "成交量不能预测一切，但能帮助判断价格变化是否有资金参与。", page)
    add_card(slide, 0.72, 1.28, 12.0, 5.45, rgb("F7FAFC"))
    x0, y0, w, h = 1.25, 1.78, 6.2, 4.55
    # axes
    ax1 = slide.shapes.add_connector(1, Inches(x0), Inches(y0 + h / 2), Inches(x0 + w), Inches(y0 + h / 2))
    ax2 = slide.shapes.add_connector(1, Inches(x0 + w / 2), Inches(y0), Inches(x0 + w / 2), Inches(y0 + h))
    for ax in (ax1, ax2):
        ax.line.color.rgb = rgb("9EB2C4")
        ax.line.width = Pt(1.5)
    labels = [
        ("放量上涨", "突破有效性提高，但要防高位诱多", x0 + 3.35, y0 + 0.45, RED),
        ("缩量上涨", "上涨动能可能不足，适合观察持续性", x0 + 0.35, y0 + 0.45, GOLD),
        ("放量下跌", "卖压明显，先控制风险", x0 + 3.35, y0 + 2.75, GREEN),
        ("缩量下跌", "恐慌减弱或无人接力，需看位置", x0 + 0.35, y0 + 2.75, BLUE),
    ]
    for title, desc, x, y, col in labels:
        add_textbox(slide, x, y, 2.3, 0.25, title, 17, col, True)
        add_textbox(slide, x, y + 0.38, 2.35, 0.75, desc, 12.5, TEXT)
    add_textbox(slide, x0 + 2.45, y0 - 0.3, 1.5, 0.25, "价格上涨", 10.5, MUTED, False, PP_ALIGN.CENTER)
    add_textbox(slide, x0 + 2.45, y0 + h + 0.06, 1.5, 0.25, "价格下跌", 10.5, MUTED, False, PP_ALIGN.CENTER)
    add_textbox(slide, x0 - 0.2, y0 + 2.03, 1.4, 0.25, "缩量", 10.5, MUTED)
    add_textbox(slide, x0 + w - 0.62, y0 + 2.03, 1.4, 0.25, "放量", 10.5, MUTED)
    add_card(slide, 8.1, 1.82, 3.9, 4.15, rgb("EEF4F8"))
    add_textbox(slide, 8.45, 2.12, 3.2, 0.28, "新手读量的三个问题", 17, TEXT, True)
    add_bullets(slide, 8.48, 2.72, 3.1, 2.6, [
        "今天的量比过去一段时间大还是小？",
        "放量发生在突破、下跌还是高位震荡？",
        "量能变化是否支持你的交易计划？",
    ], 14)
    return slide


def flow_slide(page):
    slide = base_slide("买入前检查清单：先写计划，再按计划交易", "交易框架", "交易纪律的价值，是让你在情绪最强时仍有一套可执行流程。", page)
    add_card(slide, 0.72, 1.28, 12.0, 5.45, rgb("F7FAFC"))
    steps = [
        ("看规则", "市场/交易单位/费用"),
        ("看趋势", "指数与板块环境"),
        ("看位置", "支撑、压力、均线"),
        ("定风险", "止损点与仓位"),
        ("写复盘", "结果与错误类型"),
    ]
    y = 2.25
    for i, (t, d) in enumerate(steps):
        x = 0.95 + i * 2.35
        box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(1.95), Inches(1.18))
        box.adjustments[0] = 0.10
        set_fill(box, [BLUE, CYAN, GOLD, RED, GREEN][i])
        no_line(box)
        add_textbox(slide, x + 0.12, y + 0.22, 1.7, 0.25, t, 17, WHITE, True, PP_ALIGN.CENTER)
        add_textbox(slide, x + 0.12, y + 0.62, 1.7, 0.28, d, 9.5, WHITE, False, PP_ALIGN.CENTER)
        if i < len(steps) - 1:
            arrow = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RIGHT_ARROW, Inches(x + 1.98), Inches(y + 0.41), Inches(0.32), Inches(0.34))
            set_fill(arrow, rgb("9EB2C4"))
            no_line(arrow)
    add_textbox(slide, 1.0, 4.2, 10.8, 0.62, "核心原则：每一次下单，都要能回答“为什么买、错了怎么办、赚了怎么卖”。", 22, TEXT, True, PP_ALIGN.CENTER)
    return slide


STOCK_CASES = [
    {"market": "A股", "name": "贵州茅台", "code": "600519.SH", "tag": "消费龙头", "color": RED,
     "points": ["高价股", "品牌/利润率", "适合讲长期趋势"], "series": [2.4, 2.6, 3.0, 3.8, 4.2, 4.0, 4.7, 4.4]},
    {"market": "A股", "name": "宁德时代", "code": "300750.SZ", "tag": "新能源成长", "color": GREEN,
     "points": ["成长股", "产业景气", "适合讲波动"], "series": [1.8, 2.5, 3.4, 4.7, 4.0, 3.2, 3.6, 3.1]},
    {"market": "A股", "name": "工商银行", "code": "601398.SH", "tag": "大型银行", "color": BLUE,
     "points": ["低波动", "分红/估值", "适合讲防守"], "series": [2.6, 2.55, 2.7, 2.65, 2.78, 2.9, 2.85, 3.0]},
    {"market": "港股", "name": "腾讯控股", "code": "00700.HK", "tag": "互联网平台", "color": CYAN,
     "points": ["港币计价", "平台经济", "适合讲估值回撤"], "series": [3.2, 4.0, 4.6, 3.8, 2.8, 3.2, 3.6, 3.4]},
    {"market": "港股", "name": "阿里巴巴-W", "code": "09988.HK", "tag": "电商/云", "color": GOLD,
     "points": ["双重上市", "电商竞争", "适合讲事件影响"], "series": [3.8, 3.4, 2.8, 2.3, 2.6, 2.2, 2.7, 2.5]},
    {"market": "港股", "name": "汇丰控股", "code": "00005.HK", "tag": "国际银行", "color": BLUE,
     "points": ["利率敏感", "高股息关注", "适合讲宏观影响"], "series": [2.4, 2.3, 2.6, 2.9, 3.1, 3.4, 3.2, 3.5]},
    {"market": "美股", "name": "Apple", "code": "AAPL", "tag": "消费电子", "color": rgb("5A6B7D"),
     "points": ["全球龙头", "回购/生态", "适合讲大盘权重"], "series": [2.2, 2.5, 2.8, 3.0, 3.4, 3.2, 3.8, 4.1]},
    {"market": "美股", "name": "NVIDIA", "code": "NVDA", "tag": "AI芯片", "color": RED,
     "points": ["高成长", "景气周期", "适合讲趋势加速"], "series": [1.5, 1.8, 2.4, 3.2, 4.4, 5.2, 5.8, 6.4]},
    {"market": "美股", "name": "Tesla", "code": "TSLA", "tag": "高波动成长", "color": GREEN,
     "points": ["波动大", "叙事强", "适合讲仓位控制"], "series": [2.0, 3.4, 2.7, 4.5, 3.1, 4.0, 2.6, 3.3]},
]


def draw_mini_line(slide, x, y, w, h, vals, col, fill=False):
    mn, mx = min(vals), max(vals)
    if mx == mn:
        mx += 1
    pts = []
    for i, v in enumerate(vals):
        px = x + i * w / (len(vals) - 1)
        py = y + h - (v - mn) / (mx - mn) * h
        pts.append((px, py))
    for i in range(len(pts) - 1):
        line = slide.shapes.add_connector(1, Inches(pts[i][0]), Inches(pts[i][1]), Inches(pts[i + 1][0]), Inches(pts[i + 1][1]))
        line.line.color.rgb = col
        line.line.width = Pt(2.1)
    if fill:
        base_y = y + h
        for i, (px, py) in enumerate(pts):
            bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(px - 0.035), Inches(py), Inches(0.07), Inches(base_y - py))
            set_fill(bar, col, 55)
            no_line(bar)


def stock_case_grid_slide(page):
    slide = base_slide("真实股票案例池：用来练规则和图表，不用来荐股", "案例图解", "这些股票只作为教学样本；小走势图为教学示意，非真实历史价格。", page)
    add_card(slide, 0.55, 1.18, 12.25, 5.68, rgb("F7FAFC"))
    for idx, case in enumerate(STOCK_CASES):
        row, col = divmod(idx, 3)
        x = 0.82 + col * 4.0
        y = 1.48 + row * 1.72
        card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(3.68), Inches(1.38))
        card.adjustments[0] = 0.06
        set_fill(card, WHITE)
        card.line.color.rgb = rgb("D7E1EA")
        band = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(0.16), Inches(1.38))
        set_fill(band, case["color"])
        no_line(band)
        add_textbox(slide, x + 0.32, y + 0.16, 1.85, 0.26, case["name"], 15, TEXT, True, margin=0)
        add_textbox(slide, x + 2.28, y + 0.18, 1.05, 0.22, case["code"], 9.5, MUTED, False, PP_ALIGN.RIGHT, margin=0)
        add_textbox(slide, x + 0.32, y + 0.50, 1.35, 0.24, case["market"] + " · " + case["tag"], 9.5, case["color"], True, margin=0)
        add_textbox(slide, x + 0.32, y + 0.83, 1.5, 0.34, " / ".join(case["points"][:2]), 9.5, MUTED, margin=0)
        draw_mini_line(slide, x + 2.05, y + 0.64, 1.18, 0.43, case["series"], case["color"])
    return slide


def stock_map_slide(page):
    slide = base_slide("把股票放进坐标系：先理解类型，再看K线", "案例图解", "同样一根K线，放在银行股、成长股和互联网平台股上，解释会不同。", page)
    add_card(slide, 0.72, 1.25, 12.0, 5.45, rgb("F7FAFC"))
    x0, y0, w, h = 1.35, 1.9, 9.6, 3.95
    axh = slide.shapes.add_connector(1, Inches(x0), Inches(y0 + h), Inches(x0 + w), Inches(y0 + h))
    axv = slide.shapes.add_connector(1, Inches(x0), Inches(y0 + h), Inches(x0), Inches(y0))
    for ax in (axh, axv):
        ax.line.color.rgb = rgb("A7B8C8")
        ax.line.width = Pt(1.7)
    add_textbox(slide, x0 + 3.65, y0 + h + 0.18, 2.2, 0.24, "成长/波动更高", 11, MUTED, False, PP_ALIGN.CENTER)
    add_textbox(slide, x0 - 0.22, y0 - 0.28, 2.3, 0.24, "商业稳定性更高", 11, MUTED)
    coords = {
        "贵州茅台": (4.0, 1.45), "宁德时代": (6.8, 2.25), "工商银行": (2.2, 3.0),
        "腾讯控股": (5.3, 2.1), "阿里巴巴-W": (5.8, 2.75), "汇丰控股": (2.75, 2.65),
        "Apple": (4.4, 1.65), "NVIDIA": (8.3, 1.65), "Tesla": (8.0, 3.15),
    }
    by_name = {c["name"]: c for c in STOCK_CASES}
    for name, (px, py) in coords.items():
        c = by_name[name]
        dot = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x0 + px), Inches(y0 + py), Inches(0.18), Inches(0.18))
        set_fill(dot, c["color"])
        no_line(dot)
        add_textbox(slide, x0 + px + 0.18, y0 + py - 0.05, 1.3, 0.22, name, 8.5, TEXT, True, margin=0)
    add_card(slide, 9.95, 2.25, 2.18, 2.55, rgb("EEF4F8"))
    add_textbox(slide, 10.18, 2.54, 1.7, 0.26, "读图顺序", 15, TEXT, True, PP_ALIGN.CENTER)
    add_bullets(slide, 10.15, 3.0, 1.7, 1.35, ["先看市场规则", "再看行业属性", "最后看图形信号"], 11.5, gap=4)
    return slide


def stock_visual_case_slide(page, title, cases, caption):
    slide = base_slide(title, "案例图解", caption + " 小走势图为教学示意，非真实历史价格。", page)
    for idx, case in enumerate(cases):
        x = 0.72 + idx * 4.08
        y = 1.35
        add_card(slide, x, y, 3.7, 5.25, rgb("F7FAFC"))
        add_textbox(slide, x + 0.28, y + 0.28, 2.2, 0.34, case["name"], 21, TEXT, True, margin=0)
        add_textbox(slide, x + 2.48, y + 0.34, 0.85, 0.2, case["code"], 9, MUTED, False, PP_ALIGN.RIGHT, margin=0)
        badge = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x + 0.28), Inches(y + 0.82), Inches(1.35), Inches(0.32))
        set_fill(badge, case["color"])
        no_line(badge)
        add_textbox(slide, x + 0.28, y + 0.87, 1.35, 0.15, case["tag"], 8.5, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 0)
        draw_mini_line(slide, x + 0.35, y + 1.58, 2.95, 1.35, case["series"], case["color"], True)
        add_textbox(slide, x + 0.32, y + 3.25, 2.95, 0.24, "适合练习", 12.5, MUTED, True, PP_ALIGN.CENTER)
        add_bullets(slide, x + 0.50, y + 3.66, 2.6, 1.05, case["points"], 12.5, gap=2)
    return slide


def stock_kline_context_slide(page):
    slide = base_slide("同样看K线，不同股票要问不同问题", "案例图解", "K线形态只是入口；本页K线为教学示意，非真实历史价格。", page)
    add_card(slide, 0.72, 1.20, 12.0, 5.55, rgb("F7FAFC"))
    examples = [
        ("贵州茅台", "消费龙头", RED, "更关注长期趋势与估值消化"),
        ("腾讯控股", "港股互联网", CYAN, "同时看港股流动性和政策/业绩预期"),
        ("Tesla", "高波动成长", GREEN, "仓位和止损比形态名字更重要"),
    ]
    for i, (name, tag, col, tip) in enumerate(examples):
        x = 1.05 + i * 3.95
        add_textbox(slide, x, 1.58, 2.0, 0.28, name, 18, TEXT, True)
        add_textbox(slide, x, 1.92, 2.0, 0.22, tag, 10.5, col, True)
        vals = [
            (72, 84, 88, 70, ""), (84, 78, 88, 75, ""), (78, 91, 94, 76, ""),
            (91, 88, 95, 85, ""), (88, 94, 97, 86, "")
        ] if i == 0 else [
            (86, 73, 90, 70, ""), (73, 82, 85, 67, ""), (82, 78, 88, 75, ""),
            (78, 91, 93, 73, ""), (91, 84, 96, 81, "")
        ] if i == 1 else [
            (70, 92, 96, 65, ""), (92, 68, 98, 63, ""), (68, 87, 92, 66, ""),
            (87, 73, 90, 69, ""), (73, 95, 99, 70, "")
        ]
        for j, item in enumerate(vals):
            draw_candle(slide, x + 0.55 + j * 0.43, 2.28, item[0], item[1], item[2], item[3], scale=0.035)
        add_textbox(slide, x, 5.10, 3.15, 0.62, tip, 14, TEXT, False, PP_ALIGN.CENTER)
    return slide


def checklist_slide(page):
    headers = ["项目", "记录内容", "示例问题"]
    rows = [
        ["买入理由", "趋势、形态、基本面、事件", "这次买入是不是有明确依据？"],
        ["风险设定", "止损位、仓位、最大亏损", "若判断错误，亏损是否可承受？"],
        ["卖出理由", "止盈、止损、计划变化", "卖出是按计划，还是被情绪推着走？"],
        ["结果归因", "市场、板块、个股、执行", "这次结果和能力有关，还是运气因素？"],
        ["下次改进", "保留/删除/调整的规则", "下一笔交易具体要改变什么？"],
    ]
    return table_slide("新手复盘表：把感觉变成证据", "交易框架", headers, rows, "复盘不是责备自己，而是把模糊经验沉淀成下一次能用的规则。", page)


def make_deck():
    page = 1
    cover_slide(page); page += 1
    section_slide("这套课要解决什么", "先把市场规则、图表语言和风控框架搭起来。", [
        "看懂 A股、港股、美股的基础交易规则",
        "能解释一根 K 线、趋势、均线和成交量",
        "知道技术指标的作用与局限",
        "建立买入前计划、仓位控制和复盘习惯",
    ], page); page += 1
    section_slide("目录", "建议用 60-90 分钟讲完，最后留 15 分钟做图表练习。", [
        "第一部分：股票和交易规则",
        "第二部分：A股、港股、美股对比",
        "第三部分：K线、均线、成交量和指标",
        "第四部分：交易计划、风险控制和复盘",
    ], page); page += 1

    slides = [
        ("股票到底是什么", "股票基础", [
            "股票代表对一家公司的部分所有权，价格由买卖双方在市场中形成。",
            "股价上涨不等于公司一定变好，下跌也不等于公司一定变坏。",
            "新手先区分：公司价值、市场情绪、资金流动、交易规则。",
            "买股票不是买故事，而是在不确定性里承担风险并寻求回报。",
        ], "先让新手从“赌涨跌”切换到“理解资产和风险”。", "四个关键词", ["所有权", "价格", "风险", "规则"]),
        ("开户后你能做什么", "股票基础", [
            "证券账户用于持有股票，资金账户用于买卖结算。",
            "A股、港股、美股通常需要不同权限或通道，费用和币种也不同。",
            "可操作的不止买卖股票，还包括申购、分红到账、查询公告等。",
            "不同券商的界面不同，但底层流程基本都是委托、成交、清算、交收。",
        ], "不要假设开户等于会交易，账户权限和市场入口是第一层门槛。", "账户动作", ["入金", "委托", "成交", "持仓", "交收"]),
        ("一次买卖的完整流程", "股票基础", [
            "下单前：选择市场、代码、方向、价格、数量、订单类型。",
            "交易中：委托进入交易系统，符合价格优先和时间优先原则后成交。",
            "成交后：账户出现持仓和可用资金变化，但交收通常不是立即完成。",
            "复盘时：记录成交价格、滑点、手续费和是否按计划执行。",
        ], "把流程讲清楚，新手才知道自己每一步在做什么。", "流程", ["下单", "撮合", "成交", "清算", "交收"]),
        ("订单类型：市价与限价", "股票基础", [
            "限价单：指定可接受价格，优点是价格可控，缺点是可能不成交。",
            "市价单：优先成交，价格不完全可控，在波动大或流动性差时风险更高。",
            "新手更适合从限价单开始，先控制成交价格。",
            "任何订单都要先确认市场、代码、方向和数量，避免低级错误。",
        ], "市价单不是“更高级”，只是不同的取舍。", "选择逻辑", ["要价格", "用限价", "要速度", "看流动性"]),
        ("交易费用与真实收益", "股票基础", [
            "股票交易通常涉及佣金、印花税、交易规费、平台费或换汇成本。",
            "频繁交易会让小幅盈利被费用吞掉，尤其对短线新手影响明显。",
            "跨市场投资还要考虑汇率波动和换汇差价。",
            "看收益要看扣费后的净收益，而不是只看买卖价差。",
        ], "费用是新手最容易忽略、但最稳定发生的成本。", "成本来源", ["佣金", "税费", "平台费", "汇率"]),
        ("交易日、休市与公告", "股票基础", [
            "不同市场节假日不同，A股、港股、美股不一定同一天开市。",
            "财报、分红、停牌、重大公告会影响交易和价格波动。",
            "新手不应在不了解公告的情况下追涨杀跌。",
            "跨市场持仓要额外关注当地假期和夜间消息。",
        ], "休市和公告经常让新手误判“为什么不能交易”。", "检查清单", ["交易日", "停牌", "财报", "分红", "假期"]),
        ("交收：成交不等于所有事情完成", "股票基础", [
            "成交是买卖匹配成功，交收是证券和资金最终完成划转。",
            "A股、美股、港股的交收周期和可卖规则不同。",
            "T+1、T+2 中的 T 指交易日，不是自然日。",
            "理解交收规则，能避免“为什么今天买了不能卖/钱不能提现”的困惑。",
        ], "用 T 日概念先打底，后面三地规则更容易讲。", "概念区分", ["成交", "可用", "可取", "交收"]),
    ]
    for title, section, bullets, note, side_title, side_items in slides:
        text_slide(title, section, bullets, note, page, side_title, side_items); page += 1

    table_slide("三地股票市场规则总览", "交易规则",
        ["市场", "交易时段", "交收", "交易单位", "涨跌幅/波动控制", "新手关注点"],
        [
            ["A股", "常见为 9:15-9:25 集合竞价；9:30-11:30、13:00-15:00 连续竞价", "股票通常 T+1 可卖；资金可用/可取规则需看券商", "主板通常 100 股一手", "主板常见 10%；科创板/创业板常见 20%；新股阶段另有规则", "规则细、板块差异多，先确认代码所属板块"],
            ["港股", "开市前、早市、午市及收市竞价；遇假期/恶劣天气有安排", "通常 T+2", "每只股票一手股数可能不同", "无普通日涨跌停；设有市场波动调节机制 VCM", "注意港币、每手股数、流动性和费用"],
            ["美股", "核心交易时段通常为美东 9:30-16:00；另有盘前盘后", "多数证券自 2024-05-28 起 T+1", "通常 1 股起", "无普通日涨跌停；有市场熔断与个股波动暂停机制", "注意美元、时差、盘前盘后流动性"],
        ],
        "这页只讲常见股票规则，具体品种、账户权限和券商规则以最新公告为准。", page,
        "资料参考：交易所/SEC/HKEX/NYSE 官方规则与公告，口径截至 2026-06-12。"); page += 1

    stock_case_grid_slide(page); page += 1
    stock_visual_case_slide(page, "A股案例：同在A股，股票性格也完全不同", STOCK_CASES[0:3],
                            "贵州茅台、宁德时代、工商银行用于展示消费、成长、金融三类样本。"); page += 1
    stock_visual_case_slide(page, "港股案例：规则、币种和流动性要一起看", STOCK_CASES[3:6],
                            "腾讯控股、阿里巴巴-W、汇丰控股用于展示港股平台、科技和金融样本。"); page += 1
    stock_visual_case_slide(page, "美股案例：1股起更灵活，但波动和时差更要管", STOCK_CASES[6:9],
                            "Apple、NVIDIA、Tesla用于展示美股权重股、AI成长股和高波动成长股。"); page += 1
    stock_map_slide(page); page += 1

    rule_slides = [
        ("A股交易规则：先分清板块", "交易规则", [
            "沪深主板、科创板、创业板在涨跌幅、新股交易机制、投资者门槛上存在差异。",
            "常见交易日为周一至周五，法定节假日和交易所公告休市除外。",
            "常见股票交易单位为 100 股及其整数倍，卖出零股有特殊处理。",
            "新手下单前先确认：市场、板块、涨跌幅限制、是否停牌、是否有风险警示。",
        ], "A股不是一个规则完全一致的整体，代码所在板块很关键。", "板块差异", ["主板", "科创板", "创业板", "风险警示"]),
        ("A股 T+1：今天买入，通常不能当天卖出", "交易规则", [
            "A股股票常见交易制度为 T+1：T 日买入，通常 T+1 日才可卖出。",
            "资金层面还要区分可用资金与可取资金，提现通常有额外时间要求。",
            "T 指交易日，不包括周末和休市日。",
            "新手做短线时，必须把不能当天卖出的风险放进计划里。",
        ], "很多冲动交易的问题，来自忽略 T+1 的退出约束。", "容易误解", ["能买", "能卖", "可用", "可取"]),
        ("港股交易规则：没有普通涨跌停，但不是没有风控", "交易规则", [
            "港股普通股票没有类似 A股每日固定涨跌停的制度。",
            "市场设有波动调节机制 VCM，极端波动时可能触发冷静期。",
            "港股通常以港币交易，交易单位由上市公司决定，一手股数不统一。",
            "新手要特别关注价差、流动性、交易费用和汇率。",
        ], "港股“没有涨跌停”容易被误解为风险更小，实际是波动管理方式不同。", "港股关键词", ["T+2", "港币", "一手不同", "VCM"]),
        ("美股交易规则：时间长、工具多、波动也更开放", "交易规则", [
            "美股核心交易时段通常为美东 9:30-16:00，另有盘前和盘后交易。",
            "多数美股证券结算已从 2024-05-28 起进入 T+1 周期。",
            "美股通常没有普通日涨跌停，但有大盘熔断和个股波动暂停机制。",
            "新手谨慎参与盘前盘后，因为流动性和价差可能显著变化。",
        ], "盘前盘后看起来更自由，但对新手更考验价格控制。", "美股关键词", ["T+1", "美元", "时差", "熔断"]),
        ("跨市场投资的额外风险", "交易规则", [
            "汇率会影响最终收益：股价赚钱，不代表换回本币后一定赚钱。",
            "不同市场信息披露语言、财报周期、交易假期都不同。",
            "跨市场交易费用结构更复杂，频繁操作成本更高。",
            "新手优先选择自己能理解规则、费用和信息来源的市场。",
        ], "跨市场不是简单多一个账户，而是多一套规则和风险。", "额外变量", ["汇率", "时差", "税费", "信息"]),
    ]
    for title, section, bullets, note, side_title, side_items in rule_slides:
        text_slide(title, section, bullets, note, page, side_title, side_items); page += 1

    k_slides = [
        ("一根K线包含四个价格", [
            "开盘价：该周期第一笔成交附近的价格。",
            "收盘价：该周期最后成交附近的价格，常被市场最重视。",
            "最高价/最低价：该周期波动范围。",
            "K线读的是价格行为，不是单独的预测答案。",
        ], "先把四个价格讲清，再讲形态。"),
        ("阳线、阴线与影线", [
            "阳线：收盘价高于开盘价，通常表示该周期买方占优。",
            "阴线：收盘价低于开盘价，通常表示该周期卖方占优。",
            "上影线长：上方曾被拉高，但收盘未能维持。",
            "下影线长：下方曾被打低，但收盘有所收回。",
        ], "颜色习惯在不同软件可能不同，关键看开收盘关系。"),
        ("十字星：代表犹豫，不代表一定反转", [
            "十字星说明开盘价和收盘价接近，多空暂时拉锯。",
            "低位十字星、高位十字星、震荡区十字星的意义不同。",
            "必须结合成交量、前后K线和趋势位置判断。",
            "新手不要看到十字星就自动认为马上变盘。",
        ], "技术形态永远要放到位置和趋势里看。"),
        ("大阳线与大阴线：强弱信号也会失真", [
            "大阳线通常显示买盘强，但高位大阳也可能是最后冲刺。",
            "大阴线通常显示卖压强，但低位大阴也可能是恐慌释放。",
            "同一根K线，在不同市场环境下含义不同。",
            "看大K线时必须问：突破了什么？跌破了什么？成交量是否配合？",
        ], "不要只看线体长度，要看它改变了什么结构。"),
        ("位置比形态更重要", [
            "支撑附近的长下影，和高位加速后的长下影，含义可能完全不同。",
            "压力附近的大阳线，若不能放量突破，可能只是冲高回落。",
            "趋势中继形态和顶部/底部形态不能混着用。",
            "先画区间、趋势线和均线，再解释单根K线。",
        ], "这页是K线部分的重点：形态离开位置会误导新手。"),
        ("把多根K线连起来看", [
            "连续抬高的低点，通常比单日大涨更能说明趋势改善。",
            "连续降低的高点，说明反弹力度可能偏弱。",
            "震荡区内的K线信号成功率较低，容易反复打脸。",
            "新手先识别趋势、震荡、突破、回落四种状态。",
        ], "从单根K线过渡到市场结构。"),
        ("K线常见误区", [
            "误区一：把形态当成确定性预测。",
            "误区二：只看日K，不看大盘、板块和成交量。",
            "误区三：把短周期信号用于长线决策。",
            "误区四：看对方向，却没有仓位和止损计划。",
        ], "K线能帮你观察市场，但不能替你承担风险。"),
    ]
    for title, bullets, note in k_slides:
        kline_slide(title, bullets, note, page); page += 1
    stock_kline_context_slide(page); page += 1

    ma_chart_slide(page); page += 1
    ma_text = [
        ("常用均线怎么分工", "均线系统", [
            "5日、10日：偏短线，反应快，容易受情绪波动影响。",
            "20日：常被视作月度节奏线，适合观察中短期趋势。",
            "60日：常用于观察中期趋势和阶段成本。",
            "120日、250日：更偏长期视角，适合判断大方向。",
        ], "均线周期没有神秘性，本质是不同时间长度的平均成本。", "周期分工", ["短期", "中期", "长期", "大方向"]),
        ("金叉、死叉与均线排列", "均线系统", [
            "金叉：短期均线上穿长期均线，常被视为趋势改善信号。",
            "死叉：短期均线下穿长期均线，常被视为趋势转弱信号。",
            "多头排列：短中长期均线由上到下排列，说明趋势较强。",
            "空头排列：短中长期均线由下到上排列，反弹压力可能较大。",
        ], "金叉死叉是滞后信号，不能脱离位置和成交量。", "信号类型", ["金叉", "死叉", "多头", "空头"]),
        ("均线什么时候容易失效", "均线系统", [
            "横盘震荡中，均线会频繁缠绕，金叉死叉容易反复失真。",
            "突发消息或财报冲击时，价格可能直接跳空越过均线。",
            "强趋势末端，均线仍可能看起来很好，但风险已经累积。",
            "新手不要用单一均线作为唯一买卖依据。",
        ], "所有指标都有使用环境，均线也不例外。", "失效环境", ["震荡", "跳空", "末端", "单指标"]),
    ]
    for title, section, bullets, note, side_title, side_items in ma_text:
        text_slide(title, section, bullets, note, page, side_title, side_items); page += 1

    quadrant_slide(page); page += 1
    indicator_slides = [
        ("MACD：看趋势变化，不看神秘符号", "成交量与指标", [
            "MACD 常用于观察趋势动能变化，包括快慢线、柱状图和零轴。",
            "金叉、死叉只是提示趋势可能变化，不代表一定买卖。",
            "背离可以提示风险，但背离可能持续很久。",
            "MACD 更适合趋势行情，震荡行情中容易反复发出噪声。",
        ], "讲指标时要一直强调：信号不是指令。", "MACD看什么", ["趋势", "动能", "零轴", "背离"]),
        ("RSI、KDJ、布林线：辅助判断强弱和波动", "成交量与指标", [
            "RSI 常用于观察相对强弱，超买超卖不等于马上反转。",
            "KDJ 更敏感，短线提示多，但假信号也多。",
            "布林线观察价格相对波动区间，上轨下轨不是机械买卖点。",
            "指标越多不一定越好，新手先学少数几个，并知道它们的局限。",
        ], "不要把屏幕塞满指标，反而看不见价格本身。", "指标定位", ["强弱", "敏感", "波动", "辅助"]),
    ]
    for title, section, bullets, note, side_title, side_items in indicator_slides:
        text_slide(title, section, bullets, note, page, side_title, side_items); page += 1

    flow_slide(page); page += 1
    risk_slides = [
        ("仓位管理：先活下来，再谈提高收益", "交易框架", [
            "不要把一次判断押成满仓，尤其是新手还没有稳定方法时。",
            "单笔亏损上限应提前设定，例如不让一次错误伤到整体账户。",
            "加仓必须有规则，不能因为亏损后想摊低成本就随意加。",
            "现金也是仓位，等待也是策略的一部分。",
        ], "仓位管理是把错误控制在可承受范围内。", "仓位原则", ["分批", "上限", "不摊平", "留现金"]),
        ("止损和止盈：退出比买入更难", "交易框架", [
            "止损不是承认失败，而是承认市场和自己都可能出错。",
            "止损可以按价格、技术位、资金亏损比例或时间来设计。",
            "止盈可以分批，也可以根据趋势破坏来执行。",
            "没有退出规则的交易，很容易从短线变成长线被动持有。",
        ], "退出规则要在买入前写好，而不是亏损后再想。", "退出方式", ["价格", "技术", "资金", "时间"]),
        ("新手最常见的五个坑", "交易框架", [
            "追涨杀跌：情绪驱动下单，买在兴奋点，卖在恐慌点。",
            "亏损加仓：没有计划地摊低成本，扩大单一错误。",
            "只看K线：忽视规则、费用、公告、财报和市场环境。",
            "频繁换方法：每亏一次就换指标，永远无法验证方法。",
            "不复盘：重复犯错却没有记录。",
        ], "这页适合讲真实案例，但不要使用具体荐股。", "错误来源", ["情绪", "仓位", "信息", "方法", "复盘"]),
    ]
    for title, section, bullets, note, side_title, side_items in risk_slides:
        text_slide(title, section, bullets, note, page, side_title, side_items); page += 1
    checklist_slide(page); page += 1

    slide = base_slide("总结：新手先建立框架，再追求技巧", "结尾", "最后再次强调：本课件用于投资者教育，不构成任何投资建议。", page)
    add_card(slide, 0.72, 1.28, 12.0, 5.45, rgb("F7FAFC"))
    add_textbox(slide, 1.05, 1.62, 10.9, 0.45, "记住这五句话", 24, TEXT, True, PP_ALIGN.CENTER)
    add_bullets(slide, 1.55, 2.35, 10.2, 2.85, [
        "规则先于策略：不懂交易规则，就不要急着下单。",
        "K线看位置：形态离开趋势和成交量，意义会大幅下降。",
        "指标只辅助：没有任何指标能保证收益。",
        "计划先于情绪：买入前写清理由、止损、仓位和退出。",
        "风险永远第一：本课件不构成投资建议，投资需独立判断并自行承担风险。",
    ], 17)
    add_textbox(slide, 1.35, 6.05, 10.7, 0.28,
                "参考来源：上交所/深交所交易规则，HKEX 交易时段与 VCM 说明，NYSE 交易时段，SEC T+1 结算公告。口径截至 2026-06-12。",
                9, MUTED, False, PP_ALIGN.CENTER)
    page += 1

    prs.save(OUT)


if __name__ == "__main__":
    make_deck()
    print(OUT)
