from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(r"C:\Users\Singmin\Desktop\stock")
OUT = ROOT / "炒股基础知识_简洁优化版.pptx"
CHARTS = ROOT / "real_stock_charts"

FONT = "Microsoft YaHei"
BG = RGBColor(16, 28, 40)
HEADER = RGBColor(28, 45, 62)
CARD = RGBColor(248, 251, 253)
CARD2 = RGBColor(236, 244, 250)
TEXT = RGBColor(34, 47, 60)
MUTED = RGBColor(96, 110, 124)
WHITE = RGBColor(255, 255, 255)
RED = RGBColor(218, 72, 86)
GREEN = RGBColor(24, 156, 111)
BLUE = RGBColor(72, 136, 210)
GOLD = RGBColor(224, 166, 58)
CYAN = RGBColor(58, 178, 190)


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def no_line(shape):
    shape.line.fill.background()


def bg(slide):
    rect = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    fill(rect, BG)
    no_line(rect)
    top = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.92))
    fill(top, HEADER)
    no_line(top)


def text(slide, x, y, w, h, value, size=18, color=TEXT, bold=False,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = value
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def card(slide, x, y, w, h, color=CARD):
    c = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    c.adjustments[0] = 0.06
    fill(c, color)
    c.line.color.rgb = RGBColor(220, 229, 237)
    c.line.width = Pt(0.7)
    return c


def header(slide, title, page):
    text(slide, 0.55, 0.22, 9.2, 0.42, title, 22, WHITE, True)
    text(slide, 10.45, 0.30, 2.0, 0.20, "基础知识 · 简洁版", 9, RGBColor(190, 209, 224), False, PP_ALIGN.RIGHT)
    text(slide, 12.55, 7.22, 0.45, 0.18, f"{page:02d}", 8, RGBColor(158, 174, 188), False, PP_ALIGN.RIGHT)
    text(slide, 0.55, 7.16, 12.2, 0.22, "真实股票与行情图仅作教学样本，不构成投资建议。行情图来自公开历史数据，生成于 2026-06-13。", 7.5, RGBColor(184, 198, 211))


def base(title, page):
    slide = prs.slides.add_slide(blank)
    bg(slide)
    header(slide, title, page)
    return slide


def bullets(slide, x, y, w, h, items, size=19, color=TEXT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(9)
        p._p.get_or_add_pPr().set("marL", "180000")
        p._p.get_or_add_pPr().set("indent", "-105000")


def conclusion(slide, value):
    box = card(slide, 0.82, 1.28, 11.7, 0.78, RGBColor(232, 241, 248))
    box.line.color.rgb = RGBColor(197, 214, 228)
    text(slide, 1.12, 1.48, 11.1, 0.30, value, 21, TEXT, True, PP_ALIGN.CENTER)


def cover():
    slide = prs.slides.add_slide(blank)
    bg(slide)
    text(slide, 0.75, 0.72, 4.8, 0.32, "炒股入门 · 基础知识笔记版", 15, RGBColor(190, 209, 224), True)
    text(slide, 0.70, 1.42, 8.6, 0.92, "炒股基础知识", 46, WHITE, True)
    text(slide, 0.76, 2.50, 8.3, 0.46, "少文字，多图表；先懂规则，再看行情", 22, RGBColor(218, 231, 240))
    for idx, (name, img) in enumerate([
        ("A股", "600519_SH.png"),
        ("港股", "00700_HK.png"),
        ("美股", "AAPL.png"),
    ]):
        x = 7.65
        y = 1.02 + idx * 1.75
        card(slide, x, y, 4.75, 1.42, WHITE)
        text(slide, x + 0.20, y + 0.12, 0.65, 0.18, name, 9.5, MUTED, True)
        slide.shapes.add_picture(str(CHARTS / img), Inches(x + 0.72), Inches(y + 0.14), width=Inches(3.75), height=Inches(1.08))
    text(slide, 0.76, 6.28, 7.0, 0.35, "适合：零基础、刚开户、想建立交易常识的新手", 15, RGBColor(205, 219, 230))
    text(slide, 0.55, 7.16, 12.2, 0.22, "真实股票与行情图仅作教学样本，不构成投资建议。", 7.5, RGBColor(184, 198, 211))


def note_slide(title, one_liner, items, page, accent=BLUE):
    slide = base(title, page)
    conclusion(slide, one_liner)
    card(slide, 0.82, 2.35, 7.15, 3.9, WHITE)
    bullets(slide, 1.15, 2.82, 6.45, 2.8, items, 21)
    card(slide, 8.42, 2.35, 4.1, 3.9, CARD2)
    text(slide, 8.74, 2.75, 3.45, 0.28, "课堂笔记", 20, TEXT, True, PP_ALIGN.CENTER)
    mark = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(9.85), Inches(3.42), Inches(1.25), Inches(1.25))
    fill(mark, accent)
    no_line(mark)
    text(slide, 9.92, 3.77, 1.1, 0.28, "记住", 18, WHITE, True, PP_ALIGN.CENTER)
    return slide


def market_table(page):
    slide = base("A股、港股、美股：先看规则差异", page)
    conclusion(slide, "同样是买股票，市场规则不一样，交易体验就不一样。")
    headers = ["市场", "代表案例", "交易制度", "交收/提醒"]
    rows = [
        ["A股", "贵州茅台 / 宁德时代", "股票通常 T+1：当天买入，通常不能当天卖出", "板块涨跌幅不同；先看代码属于哪个板块"],
        ["港股", "腾讯控股 / 汇丰控股", "交易可日内买卖，常说 T+0", "交收通常 T+2；注意港币、价差、流动性"],
        ["美股", "Apple / NVIDIA", "交易可日内买卖，常说 T+0", "多数证券交收 T+1；账户/券商规则会限制频繁日内交易"],
    ]
    tbl = slide.shapes.add_table(4, 4, Inches(0.82), Inches(2.35), Inches(11.7), Inches(3.95)).table
    widths = [1.2, 2.9, 3.8, 3.8]
    for i, w in enumerate(widths):
        tbl.columns[i].width = Inches(w)
    for c, h in enumerate(headers):
        cell = tbl.cell(0, c)
        cell.text = h
        fill(cell, HEADER)
        for p in cell.text_frame.paragraphs:
            p.font.name = FONT
            p.font.size = Pt(13)
            p.font.bold = True
            p.font.color.rgb = WHITE
            p.alignment = PP_ALIGN.CENTER
    for r, row in enumerate(rows, 1):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = val
            fill(cell, WHITE if r % 2 else RGBColor(239, 246, 250))
            for p in cell.text_frame.paragraphs:
                p.font.name = FONT
                p.font.size = Pt(12.5)
                p.font.color.rgb = TEXT
                p.alignment = PP_ALIGN.CENTER if c == 0 else PP_ALIGN.LEFT


def chart_slide(title, one_liner, image, items, page):
    slide = base(title, page)
    conclusion(slide, one_liner)
    card(slide, 0.70, 2.25, 7.05, 4.55, WHITE)
    slide.shapes.add_picture(str(CHARTS / image), Inches(0.98), Inches(2.52), width=Inches(6.45), height=Inches(3.85))
    card(slide, 8.08, 2.25, 4.45, 4.55, CARD2)
    bullets(slide, 8.45, 2.78, 3.65, 2.85, items, 18)


def flow(page):
    slide = base("买入前：只问三个问题", page)
    conclusion(slide, "下单前能回答，亏钱时才不会完全靠情绪。")
    steps = [
        ("为什么买？", "规则 / 趋势 / 位置"),
        ("错了怎么办？", "止损位 / 最大亏损"),
        ("赚了怎么卖？", "止盈 / 分批 / 条件"),
    ]
    for i, (a, b) in enumerate(steps):
        x = 1.05 + i * 4.05
        box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.75), Inches(3.35), Inches(1.45))
        box.adjustments[0] = 0.08
        fill(box, [BLUE, RED, GREEN][i])
        no_line(box)
        text(slide, x + 0.20, 3.05, 2.9, 0.34, a, 22, WHITE, True, PP_ALIGN.CENTER)
        text(slide, x + 0.20, 3.55, 2.9, 0.24, b, 12.5, WHITE, False, PP_ALIGN.CENTER)
    text(slide, 1.35, 5.45, 10.6, 0.42, "新手最该练的不是预测，而是把每次交易变成可复盘的动作。", 20, WHITE, True, PP_ALIGN.CENTER)


def review_table(page):
    slide = base("复盘：用一张表记录交易", page)
    conclusion(slide, "复盘不是写日记，是找出自己反复犯的错误。")
    headers = ["项目", "怎么写"]
    rows = [
        ["买入理由", "我为什么买？依据是什么？"],
        ["风险计划", "止损在哪里？最多亏多少？"],
        ["卖出理由", "是按计划，还是被情绪推着走？"],
        ["结果归因", "赚亏来自市场、个股，还是执行？"],
        ["下次改进", "下一笔交易具体改什么？"],
    ]
    tbl = slide.shapes.add_table(6, 2, Inches(1.35), Inches(2.35), Inches(10.65), Inches(4.0)).table
    tbl.columns[0].width = Inches(2.25)
    tbl.columns[1].width = Inches(8.4)
    for c, h in enumerate(headers):
        cell = tbl.cell(0, c)
        cell.text = h
        fill(cell, HEADER)
        for p in cell.text_frame.paragraphs:
            p.font.name = FONT
            p.font.size = Pt(14)
            p.font.bold = True
            p.font.color.rgb = WHITE
            p.alignment = PP_ALIGN.CENTER
    for r, row in enumerate(rows, 1):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = val
            fill(cell, WHITE if r % 2 else RGBColor(239, 246, 250))
            for p in cell.text_frame.paragraphs:
                p.font.name = FONT
                p.font.size = Pt(13.5)
                p.font.color.rgb = TEXT
                p.alignment = PP_ALIGN.CENTER if c == 0 else PP_ALIGN.LEFT


def summary(page):
    slide = base("最后记住这五句话", page)
    conclusion(slide, "基础知识的目标，是让你少犯低级错。")
    card(slide, 1.15, 2.30, 11.05, 3.9, WHITE)
    bullets(slide, 1.72, 2.85, 10.0, 2.75, [
        "先懂规则，再谈技术。",
        "K线看位置，均线看趋势，成交量看参与度。",
        "真实行情只提供证据，不提供确定答案。",
        "下单前写计划，下单后做复盘。",
        "任何案例都不构成投资建议。",
    ], 20)


def build():
    page = 1
    cover(); page += 1
    note_slide("这份PPT怎么用", "每页只记一个核心点，讲课时用真实图补充。", [
        "先讲规则，再讲图表。",
        "每页不背概念，只问关键问题。",
        "真实股票只当样本，不当推荐。",
    ], page, CYAN); page += 1
    note_slide("股票是什么", "股票是公司所有权的一小部分，价格由市场交易形成。", [
        "买股票，本质是在承担不确定性。",
        "好公司不等于好买点。",
        "价格短期受情绪和资金影响很大。",
    ], page, BLUE); page += 1
    note_slide("一次交易怎么发生", "交易不是点一下买入就结束，而是委托、成交、清算、交收。", [
        "委托：你提交买卖条件。",
        "成交：系统撮合买卖双方。",
        "交收：资金和股票最终划转。",
    ], page, GOLD); page += 1
    market_table(page); page += 1
    note_slide("交易T+0 vs 交收T+N", "先分清：能不能当天卖，和钱券几天后完成交收，是两件事。", [
        "A股：股票交易通常 T+1。",
        "港股/美股：交易可日内买卖。",
        "交收：港股通常 T+2，美股多数证券 T+1。",
    ], page, GREEN); page += 1
    note_slide("费用会吃掉收益", "短线越频繁，费用影响越明显。", [
        "看收益要看扣费后的净收益。",
        "跨市场还要考虑汇率。",
        "小资金更要控制交易频率。",
    ], page, RED); page += 1
    chart_slide("真实行情：先学会读一张图", "看行情图，先找趋势、波动、成交量。", "600519_SH.png", [
        "上半部分：价格和K线。",
        "彩色线：不同周期均线。",
        "下半部分：成交量。",
    ], page); page += 1
    chart_slide("K线：一页讲清", "K线只回答：这段时间价格怎么走过。", "00700_HK.png", [
        "实体：开盘价到收盘价。",
        "影线：最高价和最低价。",
        "形态必须放在位置里看。",
    ], page); page += 1
    chart_slide("均线：看趋势，不是看魔法", "均线是平均成本，信号天然滞后。", "NVDA.png", [
        "MA5：短期节奏。",
        "MA20：中短趋势。",
        "MA60：阶段方向。",
    ], page); page += 1
    chart_slide("成交量：看有没有人参与", "放量代表参与度提高，不代表一定上涨。", "300750_SZ.png", [
        "突破时看量能是否配合。",
        "高位放量也可能是风险。",
        "量价要结合位置判断。",
    ], page); page += 1
    note_slide("常用指标怎么定位", "指标是辅助观察，不是买卖命令。", [
        "MACD：看趋势动能。",
        "RSI/KDJ：看短期强弱。",
        "布林线：看波动区间。",
    ], page, CYAN); page += 1
    flow(page); page += 1
    note_slide("仓位管理", "先决定最多亏多少，再决定买多少。", [
        "不要一上来满仓。",
        "亏损后不要随意加仓。",
        "现金也是仓位。",
    ], page, GOLD); page += 1
    note_slide("止损止盈", "退出规则要在买入前写好。", [
        "止损：价格、技术位、资金比例。",
        "止盈：分批、移动止盈、条件变化。",
        "没有退出规则，短线容易变被动长线。",
    ], page, RED); page += 1
    review_table(page); page += 1
    note_slide("新手常见错误", "亏损很多时候不是看不懂图，而是没有计划。", [
        "追涨杀跌。",
        "只看K线，不看规则和公告。",
        "不复盘，重复同一个错误。",
    ], page, GREEN); page += 1
    summary(page); page += 1
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
