from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(r"C:\Users\Singmin\Desktop\stock")
OUT = ROOT / "炒股基础知识_零基础听得懂版_完整规则K线强化版.pptx"
CHARTS = ROOT / "real_stock_charts"

FONT = "Microsoft YaHei"
BG = RGBColor(16, 28, 40)
TOP = RGBColor(28, 45, 62)
WHITE = RGBColor(255, 255, 255)
CARD = RGBColor(248, 251, 253)
CARD2 = RGBColor(235, 244, 250)
TEXT = RGBColor(31, 45, 58)
MUTED = RGBColor(91, 106, 120)
RED = RGBColor(218, 72, 86)
GREEN = RGBColor(24, 156, 111)
BLUE = RGBColor(70, 135, 210)
GOLD = RGBColor(224, 166, 58)
CYAN = RGBColor(58, 178, 190)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def no_line(shape):
    shape.line.fill.background()


def bg(slide):
    rect = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    fill(rect, BG)
    no_line(rect)
    top = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.92))
    fill(top, TOP)
    no_line(top)


def tx(slide, x, y, w, h, value, size=18, color=TEXT, bold=False,
       align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = value
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def card(slide, x, y, w, h, color=CARD):
    c = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    c.adjustments[0] = 0.06
    fill(c, color)
    c.line.color.rgb = RGBColor(219, 229, 237)
    c.line.width = Pt(0.7)
    return c


def header(slide, title, page):
    tx(slide, 0.55, 0.22, 9.8, 0.4, title, 22, WHITE, True)
    tx(slide, 10.20, 0.30, 2.25, 0.18, "完整规则 + K线版", 9, RGBColor(190, 209, 224), False, PP_ALIGN.RIGHT)
    tx(slide, 12.55, 7.22, 0.45, 0.18, f"{page:02d}", 8, RGBColor(158, 174, 188), False, PP_ALIGN.RIGHT)
    tx(slide, 0.55, 7.16, 12.2, 0.22, "真实股票和行情图只作教学样本，不构成投资建议；具体规则以交易所和券商最新说明为准。", 7.4, RGBColor(184, 198, 211))


def base(title, page):
    slide = prs.slides.add_slide(blank)
    bg(slide)
    header(slide, title, page)
    return slide


def big_line(slide, value):
    card(slide, 0.82, 1.25, 11.7, 0.84, CARD2)
    tx(slide, 1.08, 1.47, 11.2, 0.34, value, 22, TEXT, True, PP_ALIGN.CENTER)


def bullets(slide, x, y, w, h, items, size=20, color=TEXT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, item in enumerate(items[:3]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(9)
        p._p.get_or_add_pPr().set("marL", "180000")
        p._p.get_or_add_pPr().set("indent", "-105000")


def simple(title, line, items, page, accent=BLUE):
    slide = base(title, page)
    big_line(slide, line)
    card(slide, 0.85, 2.35, 7.25, 3.9, WHITE)
    bullets(slide, 1.20, 2.82, 6.45, 2.75, items, 22)
    card(slide, 8.55, 2.35, 3.75, 3.9, CARD2)
    dot = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(9.75), Inches(3.10), Inches(1.35), Inches(1.35))
    fill(dot, accent)
    no_line(dot)
    tx(slide, 9.82, 3.50, 1.2, 0.24, "一句话", 16, WHITE, True, PP_ALIGN.CENTER)
    tx(slide, 8.84, 4.82, 3.16, 0.48, line, 14.5, TEXT, True, PP_ALIGN.CENTER)


def cover():
    slide = prs.slides.add_slide(blank)
    bg(slide)
    tx(slide, 0.76, 0.78, 6.0, 0.32, "炒股入门 · 从买卖到K线", 15, RGBColor(190, 209, 224), True)
    tx(slide, 0.70, 1.42, 9.4, 0.9, "炒股基础知识", 46, WHITE, True)
    tx(slide, 0.76, 2.48, 8.8, 0.42, "完全零基础也能听懂的规则课", 22, RGBColor(218, 231, 240))
    for i, (label, img) in enumerate([("A股", "600519_SH.png"), ("港股", "00700_HK.png"), ("美股", "AAPL.png")]):
        y = 1.02 + i * 1.75
        card(slide, 7.55, y, 4.85, 1.42, WHITE)
        tx(slide, 7.75, y + 0.12, 0.65, 0.18, label, 9.5, MUTED, True)
        slide.shapes.add_picture(str(CHARTS / img), Inches(8.45), Inches(y + 0.13), width=Inches(3.68), height=Inches(1.08))
    tx(slide, 0.76, 6.28, 7.2, 0.32, "目标：会买卖、会算赚亏、懂基本规则、看懂最基础K线。", 15, RGBColor(205, 219, 230))
    tx(slide, 0.55, 7.16, 12.2, 0.22, "真实股票和行情图只作教学样本，不构成投资建议。", 7.5, RGBColor(184, 198, 211))


def math_profit(page):
    slide = base("买入、卖出、盈利：用小学数学算", page)
    big_line(slide, "低价买，高价卖，扣掉费用后剩下的才叫盈利。")
    examples = [("买入", "10元/股 × 100股", "花 1000 元", BLUE),
                ("卖出", "12元/股 × 100股", "拿回 1200 元", GREEN),
                ("盈利", "1200 - 1000", "赚 200 元（未扣费）", RED)]
    for i, (title, formula, result, color) in enumerate(examples):
        x = 0.85 + i * 4.05
        card(slide, x, 2.55, 3.35, 2.45, WHITE)
        tx(slide, x + 0.25, 2.90, 2.85, 0.35, title, 25, color, True, PP_ALIGN.CENTER)
        tx(slide, x + 0.25, 3.55, 2.85, 0.28, formula, 16, TEXT, False, PP_ALIGN.CENTER)
        tx(slide, x + 0.25, 4.16, 2.85, 0.30, result, 18, TEXT, True, PP_ALIGN.CENTER)
    tx(slide, 1.3, 5.65, 10.7, 0.42, "真实交易还要扣手续费、税费、平台费，软件里的盈利通常会自动扣掉一部分成本。", 18, WHITE, True, PP_ALIGN.CENTER)


def loss_math(page):
    slide = base("亏损：卖得比买得便宜，就亏了", page)
    big_line(slide, "亏损不是软件故障，是卖出价格低于买入成本。")
    card(slide, 1.35, 2.55, 10.65, 2.65, WHITE)
    labels = [("买入", "10元 × 100股", BLUE), ("卖出", "8元 × 100股", GREEN), ("结果", "亏 200 元", RED)]
    for i, (a, b, c) in enumerate(labels):
        x = 1.85 + i * 2.3
        tx(slide, x, 2.98, 2.4, 0.34, a, 22, c, True, PP_ALIGN.CENTER)
        tx(slide, x, 3.75, 2.4, 0.28, b, 18 if i < 2 else 20, c if i == 2 else TEXT, i == 2, PP_ALIGN.CENTER)
    tx(slide, 1.45, 5.72, 10.5, 0.35, "新手先记住：亏损不可怕，不知道最多会亏多少才可怕。", 20, WHITE, True, PP_ALIGN.CENTER)


def account(page):
    slide = base("账户里最常见的几个词", page)
    big_line(slide, "账户就像钱包 + 仓库：现金在钱包，股票在仓库。")
    terms = [("现金", "还没买股票的钱"), ("持仓", "你已经买到的股票"), ("成本价", "你买入后的平均价格"),
             ("可用", "现在能继续买的钱"), ("可取", "现在能提现的钱")]
    for i, (name, desc) in enumerate(terms):
        x = 0.95 + (i % 3) * 4.05
        y = 2.55 + (i // 3) * 1.55
        card(slide, x, y, 3.45, 1.05, WHITE)
        tx(slide, x + 0.18, y + 0.18, 0.9, 0.24, name, 18, [BLUE, GREEN, GOLD, CYAN, RED][i], True)
        tx(slide, x + 1.05, y + 0.19, 2.1, 0.24, desc, 14, TEXT)


def order(page):
    slide = base("下单：你必须填四件事", page)
    big_line(slide, "一次下单，就是告诉系统：买/卖什么、多少钱、多少股。")
    fields = [("方向", "买入 or 卖出"), ("代码", "是哪只股票"), ("价格", "你愿意的价格"), ("数量", "买/卖多少股")]
    for i, (name, desc) in enumerate(fields):
        x = 0.95 + i * 3.05
        card(slide, x, 2.65, 2.55, 2.1, WHITE)
        tx(slide, x + 0.18, 3.02, 2.15, 0.3, name, 24, [BLUE, CYAN, GOLD, GREEN][i], True, PP_ALIGN.CENTER)
        tx(slide, x + 0.18, 3.70, 2.15, 0.25, desc, 15, TEXT, False, PP_ALIGN.CENTER)
    tx(slide, 1.2, 5.65, 10.9, 0.35, "方向填错最危险：想买填成卖，想卖填成买，结果会完全相反。", 19, WHITE, True, PP_ALIGN.CENTER)


def deal(page):
    slide = base("成交：不是你点了买，就一定买得到", page)
    big_line(slide, "成交 = 你的价格和别人愿意交易的价格对上了。")
    parts = [("你想买", "出价 10 元", "有人愿意 10 元卖", BLUE),
             ("价格对上", "买方 + 卖方", "这笔交易才成立", GREEN)]
    for i, (a, b, c, color) in enumerate(parts):
        x = 0.85 + i * 4.05
        card(slide, x, 2.35, 3.55, 2.75, WHITE)
        tx(slide, x + 0.25, 2.70, 3.0, 0.32, a, 24, color, True, PP_ALIGN.CENTER)
        tx(slide, x + 0.25, 3.38, 3.0, 0.28, b, 18, TEXT, False, PP_ALIGN.CENTER)
        tx(slide, x + 0.25, 4.02, 3.0, 0.28, c, 18, TEXT, True, PP_ALIGN.CENTER)
    card(slide, 8.95, 2.35, 3.25, 2.75, CARD2)
    bullets(slide, 9.25, 2.78, 2.65, 1.85, ["下单 ≠ 成交。", "没人接，就挂着。", "成交后才算买到/卖掉。"], 16)
    tx(slide, 1.2, 5.78, 10.9, 0.34, "交易市场不是商店，不是你想买就一定马上有货，想卖就一定马上有人接。", 18, WHITE, True, PP_ALIGN.CENTER)


def order_book(page):
    slide = base("买一卖一、五档盘口：看谁在排队", page)
    big_line(slide, "买一是最高的买价，卖一是最低的卖价。")
    card(slide, 1.0, 2.25, 4.9, 3.65, WHITE)
    tx(slide, 1.35, 2.55, 4.2, 0.28, "卖盘：有人想卖", 19, GREEN, True, PP_ALIGN.CENTER)
    asks = [("卖五", "10.05", "900股"), ("卖四", "10.04", "600股"), ("卖三", "10.03", "800股"), ("卖二", "10.02", "500股"), ("卖一", "10.01", "1200股")]
    bids = [("买一", "10.00", "1500股"), ("买二", "9.99", "700股"), ("买三", "9.98", "900股"), ("买四", "9.97", "400股"), ("买五", "9.96", "1000股")]
    for i, row in enumerate(asks):
        y = 3.05 + i * 0.28
        tx(slide, 1.35, y, 1.0, 0.18, row[0], 10.5, GREEN, True)
        tx(slide, 2.65, y, 1.0, 0.18, row[1], 10.5, TEXT, False, PP_ALIGN.CENTER)
        tx(slide, 4.0, y, 1.0, 0.18, row[2], 10.5, MUTED, False, PP_ALIGN.RIGHT)
    tx(slide, 1.35, 4.62, 4.2, 0.24, "买盘：有人想买", 19, RED, True, PP_ALIGN.CENTER)
    for i, row in enumerate(bids):
        y = 5.08 + i * 0.28
        tx(slide, 1.35, y, 1.0, 0.18, row[0], 10.5, RED, True)
        tx(slide, 2.65, y, 1.0, 0.18, row[1], 10.5, TEXT, False, PP_ALIGN.CENTER)
        tx(slide, 4.0, y, 1.0, 0.18, row[2], 10.5, MUTED, False, PP_ALIGN.RIGHT)
    card(slide, 6.45, 2.25, 5.65, 3.65, CARD2)
    bullets(slide, 6.85, 2.85, 4.85, 2.2, ["你想马上买，通常看卖一。", "你想马上卖，通常看买一。", "买一卖一差得越远，成交成本可能越高。"], 18)


def priority(page):
    slide = base("成交优先：谁更有诚意，谁排前面", page)
    big_line(slide, "价格优先；价格一样，时间优先。")
    card(slide, 0.95, 2.45, 3.55, 2.7, WHITE)
    tx(slide, 1.20, 2.78, 3.0, 0.3, "买入排队", 22, RED, True, PP_ALIGN.CENTER)
    bullets(slide, 1.32, 3.35, 2.65, 1.25, ["出价 10.02 的人", "排在 10.01 前面", "因为买得更贵"], 14.5)
    card(slide, 4.90, 2.45, 3.55, 2.7, WHITE)
    tx(slide, 5.15, 2.78, 3.0, 0.3, "卖出排队", 22, GREEN, True, PP_ALIGN.CENTER)
    bullets(slide, 5.27, 3.35, 2.65, 1.25, ["卖价 9.98 的人", "排在 9.99 前面", "因为卖得更便宜"], 14.5)
    card(slide, 8.85, 2.45, 3.0, 2.7, CARD2)
    tx(slide, 9.1, 2.78, 2.5, 0.3, "价格一样", 20, BLUE, True, PP_ALIGN.CENTER)
    bullets(slide, 9.12, 3.45, 2.25, 0.95, ["谁先下单", "谁先排队"], 16)
    tx(slide, 1.2, 5.85, 10.9, 0.28, "所以新手要明白：挂单价格太保守，可能一直成交不了。", 18, WHITE, True, PP_ALIGN.CENTER)


def cancel_order(page):
    slide = base("撤单和未成交：挂着的可以撤，成交的不能反悔", page)
    big_line(slide, "撤单只对还没成交的部分有效。")
    items = [("未成交", "还在排队，可以撤单", BLUE), ("部分成交", "成交的留下，剩余的可撤", GOLD), ("全部成交", "交易已经成立，不能撤回", RED)]
    for i, (a, b, c) in enumerate(items):
        x = 0.95 + i * 4.05
        card(slide, x, 2.65, 3.35, 2.35, WHITE)
        tx(slide, x + 0.25, 3.0, 2.85, 0.32, a, 23, c, True, PP_ALIGN.CENTER)
        tx(slide, x + 0.25, 3.75, 2.85, 0.28, b, 15.5, TEXT, False, PP_ALIGN.CENTER)
    tx(slide, 1.25, 5.72, 10.8, 0.35, "一句话：撤单不是后悔药，只能撤还没真正成交的订单。", 19, WHITE, True, PP_ALIGN.CENTER)


def price_cage(page):
    slide = base("价格笼子：报价不能太离谱", page)
    big_line(slide, "价格笼子 = 系统给报价划一个合理范围，防止乌龙单。")
    card(slide, 0.85, 2.30, 5.25, 3.45, WHITE)
    tx(slide, 1.20, 2.68, 4.55, 0.30, "例子：当前附近价格 10 元", 22, TEXT, True, PP_ALIGN.CENTER)
    tx(slide, 1.20, 3.38, 4.55, 0.28, "你想 12 元买入", 19, RED, True, PP_ALIGN.CENTER)
    tx(slide, 1.20, 4.05, 4.55, 0.28, "离当前价格太远，可能被系统拦住", 17, TEXT, False, PP_ALIGN.CENTER)
    tx(slide, 1.20, 4.72, 4.55, 0.28, "目的：防止手滑报错价", 17, MUTED, False, PP_ALIGN.CENTER)
    card(slide, 6.65, 2.30, 5.55, 3.45, CARD2)
    bullets(slide, 7.02, 2.78, 4.85, 2.2, ["它不是涨停跌停。", "它限制离市场价太远的申报。", "A股连续竞价阶段常见有约 2% 报价约束，具体以交易所规则为准。"], 16)
    tx(slide, 1.2, 6.10, 10.9, 0.28, "价格笼子像护栏，不预测涨跌，只防止报单太夸张。", 17, WHITE, True, PP_ALIGN.CENTER)


def market_rules(page):
    slide = base("A股、港股、美股：交易规则不一样", page)
    big_line(slide, "重点分清：能不能当天卖，和几天后交收，是两件事。")
    headers = ["市场", "当天买卖", "交收", "新手提醒"]
    rows = [["A股", "股票通常 T+1\n当天买入通常不能当天卖", "资金/股票交收另有规则", "先看板块和涨跌幅"],
            ["港股", "可日内买卖\n常说 T+0", "通常 T+2", "一手股数不统一"],
            ["美股", "可日内买卖\n常说 T+0", "多数证券 T+1", "频繁日内交易看账户规则"]]
    tbl = slide.shapes.add_table(4, 4, Inches(0.78), Inches(2.45), Inches(11.8), Inches(3.55)).table
    for i, w in enumerate([1.15, 3.2, 2.5, 4.95]):
        tbl.columns[i].width = Inches(w)
    for c, h in enumerate(headers):
        cell = tbl.cell(0, c); cell.text = h; fill(cell, TOP)
        for p in cell.text_frame.paragraphs:
            p.font.name = FONT; p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = WHITE; p.alignment = PP_ALIGN.CENTER
    for r, row in enumerate(rows, 1):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c); cell.text = val; fill(cell, WHITE if r % 2 else RGBColor(240, 247, 251))
            for p in cell.text_frame.paragraphs:
                p.font.name = FONT; p.font.size = Pt(12.5); p.font.color.rgb = TEXT; p.alignment = PP_ALIGN.CENTER if c < 3 else PP_ALIGN.LEFT


def trading_phases(page):
    slide = base("交易时间和交易阶段：不是全天都一样", page)
    big_line(slide, "开盘前、盘中、收盘前，交易规则和价格形成方式可能不同。")
    rows = [("A股", "集合竞价：开盘前集中撮合\n连续竞价：盘中实时撮合"),
            ("港股", "开市前、持续交易、收市竞价\n另有午间休市和假期安排"),
            ("美股", "核心时段：美东 9:30-16:00\n盘前盘后可交易，但流动性可能更差")]
    for i, (mkt, desc) in enumerate(rows):
        x = 0.95 + i * 4.05
        card(slide, x, 2.35, 3.45, 3.15, WHITE)
        tx(slide, x + 0.20, 2.78, 3.05, 0.34, mkt, 24, [RED, CYAN, BLUE][i], True, PP_ALIGN.CENTER)
        tx(slide, x + 0.32, 3.58, 2.8, 0.92, desc, 14.5, TEXT, False, PP_ALIGN.CENTER)
    tx(slide, 1.2, 6.05, 10.9, 0.28, "新手重点：盘前盘后、收盘竞价不是普通盘中，别随便用市价单。", 17, WHITE, True, PP_ALIGN.CENTER)


def price_limit(page):
    slide = base("A股涨停、跌停：一天最多涨跌多少", page)
    big_line(slide, "涨停/跌停 = 当天价格涨跌到上限或下限。")
    card(slide, 0.85, 2.25, 3.55, 3.25, WHITE)
    tx(slide, 1.05, 2.60, 3.1, 0.30, "昨天收盘 10 元", 20, TEXT, True, PP_ALIGN.CENTER)
    tx(slide, 1.05, 3.35, 3.1, 0.28, "涨停价约 11 元", 20, RED, True, PP_ALIGN.CENTER)
    tx(slide, 1.05, 4.10, 3.1, 0.28, "跌停价约 9 元", 20, GREEN, True, PP_ALIGN.CENTER)
    card(slide, 4.90, 2.25, 3.55, 3.25, WHITE)
    tx(slide, 5.15, 2.60, 3.0, 0.30, "涨停不等于稳赚", 20, RED, True, PP_ALIGN.CENTER)
    tx(slide, 5.15, 3.35, 3.0, 0.28, "可能买不到", 17, TEXT, False, PP_ALIGN.CENTER)
    tx(slide, 5.15, 4.10, 3.0, 0.28, "第二天也可能低开", 17, TEXT, False, PP_ALIGN.CENTER)
    card(slide, 8.95, 2.25, 3.25, 3.25, CARD2)
    bullets(slide, 9.25, 2.70, 2.65, 2.05, ["主板常见 ±10%。", "ST 常见 ±5%。", "创业板/科创板常见 ±20%。"], 14.5)
    tx(slide, 1.2, 6.05, 10.9, 0.28, "涨跌停是风控规则，不是买卖理由；新股等特殊情况另算。", 17, WHITE, True, PP_ALIGN.CENTER)


def breakers(page):
    slide = base("港股、美股：没有A股式涨跌停，但有波动控制", page)
    big_line(slide, "港股/美股更开放，所以更要理解冷静期和熔断。")
    card(slide, 0.85, 2.25, 5.45, 3.55, WHITE)
    tx(slide, 1.15, 2.60, 4.85, 0.30, "港股：VCM 波动调节机制", 19, CYAN, True, PP_ALIGN.CENTER)
    bullets(slide, 1.28, 3.18, 4.5, 1.65, ["不是A股式每日涨跌停。", "部分股票短时间剧烈波动，会进入冷静期。", "冷静期内通常限制价格范围，但市场仍可交易。"], 14.5)
    card(slide, 6.85, 2.25, 5.35, 3.55, WHITE)
    tx(slide, 7.15, 2.60, 4.75, 0.30, "美股：大盘熔断 + 个股暂停", 19, GOLD, True, PP_ALIGN.CENTER)
    bullets(slide, 7.28, 3.18, 4.45, 1.65, ["大盘熔断看 S&P 500 跌幅。", "常见触发线：7%、13%、20%。", "个股也可能因波动过大临时暂停。"], 14.5)
    tx(slide, 1.2, 6.10, 10.9, 0.28, "熔断不是救市按钮，只是让市场暂停一下，给大家重新报价的时间。", 16.5, WHITE, True, PP_ALIGN.CENTER)


def halt_st(page):
    slide = base("停牌、复牌、ST：看到这些先停一下", page)
    big_line(slide, "有些股票不是随时都能买卖，有些股票风险标签更高。")
    card(slide, 0.95, 2.40, 3.35, 2.55, WHITE)
    tx(slide, 1.20, 2.75, 2.85, 0.30, "停牌", 23, BLUE, True, PP_ALIGN.CENTER)
    tx(slide, 1.20, 3.55, 2.85, 0.35, "暂时不能交易", 17, TEXT, False, PP_ALIGN.CENTER)
    card(slide, 4.95, 2.40, 3.35, 2.55, WHITE)
    tx(slide, 5.20, 2.75, 2.85, 0.30, "复牌", 23, GREEN, True, PP_ALIGN.CENTER)
    tx(slide, 5.20, 3.55, 2.85, 0.35, "恢复交易，波动可能变大", 17, TEXT, False, PP_ALIGN.CENTER)
    card(slide, 8.95, 2.40, 3.35, 2.55, WHITE)
    tx(slide, 9.20, 2.75, 2.85, 0.30, "ST/风险警示", 23, RED, True, PP_ALIGN.CENTER)
    tx(slide, 9.20, 3.55, 2.85, 0.35, "公司风险较高，涨跌幅也可能不同", 15.5, TEXT, False, PP_ALIGN.CENTER)
    tx(slide, 1.2, 5.85, 10.9, 0.28, "新手看法：有风险标签、停复牌公告的股票，先别急着碰。", 18, WHITE, True, PP_ALIGN.CENTER)


def dividend(page):
    slide = base("分红、除权除息：不是天上掉钱", page)
    big_line(slide, "分红会给你钱，但股价通常也会做相应调整。")
    card(slide, 0.90, 2.35, 5.25, 3.15, WHITE)
    tx(slide, 1.20, 2.75, 4.6, 0.28, "简单例子", 22, TEXT, True, PP_ALIGN.CENTER)
    bullets(slide, 1.35, 3.35, 4.25, 1.25, ["股价 10 元。", "每股分红 1 元。", "除息后参考价可能变成约 9 元。"], 16.5)
    card(slide, 6.70, 2.35, 5.35, 3.15, CARD2)
    bullets(slide, 7.05, 3.0, 4.55, 1.45, ["分红不是白捡。", "重点看公司长期赚钱能力。", "不要只因为“要分红”就买。"], 17)
    tx(slide, 1.2, 6.00, 10.9, 0.28, "新手看法：分红是现金流，不是保证赚钱；股价和税费也要一起看。", 17, WHITE, True, PP_ALIGN.CENTER)


def chart_slide(title, line, image, items, page):
    slide = base(title, page)
    big_line(slide, line)
    card(slide, 0.72, 2.25, 7.0, 4.55, WHITE)
    slide.shapes.add_picture(str(CHARTS / image), Inches(1.0), Inches(2.52), width=Inches(6.42), height=Inches(3.82))
    card(slide, 8.08, 2.25, 4.45, 4.55, CARD2)
    bullets(slide, 8.45, 2.78, 3.65, 2.85, items, 18)


def draw_candle(slide, x, y, up=True):
    color = RED if up else GREEN
    wick = slide.shapes.add_connector(1, Inches(x + 0.55), Inches(y), Inches(x + 0.55), Inches(y + 3.25))
    wick.line.color.rgb = color
    wick.line.width = Pt(2.0)
    body_y = y + 0.82 if up else y + 1.55
    body = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x + 0.20), Inches(body_y), Inches(0.70), Inches(1.35))
    fill(body, color)
    body.line.color.rgb = color
    return body


def draw_custom_candle(slide, x, y, open_y, close_y, high_y, low_y, label, color=None):
    if color is None:
        color = RED if close_y < open_y else GREEN
    center_x = x + 0.35
    wick = slide.shapes.add_connector(1, Inches(center_x), Inches(y + high_y), Inches(center_x), Inches(y + low_y))
    wick.line.color.rgb = color
    wick.line.width = Pt(2.0)
    top = min(open_y, close_y)
    bottom = max(open_y, close_y)
    height = max(0.06, bottom - top)
    body = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x + 0.10), Inches(y + top), Inches(0.50), Inches(height))
    fill(body, color if height > 0.08 else WHITE)
    body.line.color.rgb = color
    body.line.width = Pt(1.2)
    tx(slide, x - 0.18, y + 3.20, 1.05, 0.22, label, 9.5, TEXT, True, PP_ALIGN.CENTER)


def kline_gallery(page):
    slide = base("K线常见形态：先认形状，不急着预测", page)
    big_line(slide, "K线形态只是观察工具，不是涨跌保证。")
    card(slide, 0.65, 2.15, 12.05, 4.62, WHITE)
    samples = [
        ("大阳线", 2.45, 1.00, 0.55, 2.80, RED),
        ("大阴线", 0.70, 2.45, 0.50, 2.85, GREEN),
        ("长上影", 2.00, 2.35, 0.35, 2.65, RED),
        ("长下影", 1.35, 1.05, 0.90, 3.05, GREEN),
        ("十字星", 1.62, 1.68, 0.55, 2.95, GOLD),
        ("小阳线", 1.75, 1.40, 1.05, 2.30, RED),
        ("小阴线", 1.35, 1.75, 0.95, 2.35, GREEN),
        ("光头阳", 2.55, 0.78, 0.78, 2.78, RED),
        ("光脚阴", 0.78, 2.55, 0.55, 2.55, GREEN),
    ]
    for i, (label, open_y, close_y, high_y, low_y, color) in enumerate(samples):
        x = 0.98 + i * 1.30
        draw_custom_candle(slide, x, 2.72, open_y, close_y, high_y, low_y, label, color)
    tx(slide, 1.0, 6.20, 11.2, 0.25, "红绿颜色在不同软件里可能设置相反，真正要看的是开盘价、收盘价、最高价、最低价。", 13.5, MUTED, False, PP_ALIGN.CENTER)


def kline_location(page):
    slide = base("同一个K线形态，位置不同，意思可能完全不同", page)
    big_line(slide, "不要背形态名字，要先问：它出现在高位、低位，还是震荡中？")
    labels = [
        ("低位长下影", "可能说明下方有人接，但还要看后续确认", GREEN),
        ("高位长上影", "可能说明上方卖压大，追高要小心", RED),
        ("震荡十字星", "可能只是犹豫，不一定马上变盘", GOLD),
    ]
    for i, (title, note, color) in enumerate(labels):
        x = 0.88 + i * 4.10
        card(slide, x, 2.30, 3.55, 3.65, WHITE)
        tx(slide, x + 0.22, 2.62, 3.05, 0.30, title, 19, color, True, PP_ALIGN.CENTER)
        # Mini context line
        points = [(x + 0.45, 4.25), (x + 1.05, 4.05), (x + 1.65, 4.15), (x + 2.20, 3.75), (x + 2.85, 3.95)]
        if i == 0:
            points = [(x + 0.45, 4.65), (x + 1.05, 4.40), (x + 1.65, 4.20), (x + 2.20, 4.05), (x + 2.85, 3.70)]
            draw_custom_candle(slide, x + 1.55, 3.05, 1.85, 1.25, 1.10, 2.55, "", GREEN)
        elif i == 1:
            points = [(x + 0.45, 3.85), (x + 1.05, 3.55), (x + 1.65, 3.35), (x + 2.20, 3.20), (x + 2.85, 3.55)]
            draw_custom_candle(slide, x + 1.55, 3.05, 1.10, 1.55, 0.30, 1.75, "", RED)
        else:
            points = [(x + 0.45, 4.05), (x + 1.05, 3.90), (x + 1.65, 4.05), (x + 2.20, 3.92), (x + 2.85, 4.02)]
            draw_custom_candle(slide, x + 1.55, 3.05, 1.45, 1.50, 0.65, 2.25, "", GOLD)
        for j in range(len(points) - 1):
            line = slide.shapes.add_connector(1, Inches(points[j][0]), Inches(points[j][1]), Inches(points[j + 1][0]), Inches(points[j + 1][1]))
            line.line.color.rgb = RGBColor(120, 140, 156)
            line.line.width = Pt(1.4)
        tx(slide, x + 0.35, 5.30, 2.85, 0.34, note, 12.5, TEXT, False, PP_ALIGN.CENTER)


def kline_single(page):
    slide = base("一根K线怎么读：开盘、收盘、最高、最低", page)
    big_line(slide, "一根K线只说明这一段时间价格怎么走过。")
    card(slide, 0.72, 2.18, 6.7, 4.55, WHITE)
    slide.shapes.add_picture(str(CHARTS / "00700_HK.png"), Inches(0.95), Inches(2.48), width=Inches(6.25), height=Inches(3.85))
    card(slide, 7.75, 2.18, 4.85, 4.55, CARD2)
    draw_candle(slide, 9.60, 2.75, True)
    labels = [("最高价", 9.15, 2.66), ("收盘价", 10.45, 3.58), ("开盘价", 8.25, 4.46), ("最低价", 9.15, 5.90)]
    for label, x, y in labels:
        tx(slide, x, y, 1.2, 0.22, label, 12, TEXT, True, PP_ALIGN.CENTER)
    bullets(slide, 8.12, 5.35, 4.1, 0.75, ["实体看开盘和收盘。", "影线看最高和最低。", "不要只靠一根K线判断明天。"], 12.5)


def kline_context(page):
    slide = base("真实K线不要单独看：还要看趋势、均线、成交量", page)
    big_line(slide, "看K线的顺序：先看位置，再看形态，最后看量。")
    card(slide, 0.72, 2.18, 7.15, 4.55, WHITE)
    slide.shapes.add_picture(str(CHARTS / "600519_SH.png"), Inches(0.98), Inches(2.46), width=Inches(6.62), height=Inches(3.86))
    callouts = [("价格区", 1.15, 2.55, BLUE), ("均线", 4.45, 3.02, GOLD), ("成交量", 4.85, 5.78, CYAN)]
    for label, x, y, color in callouts:
        box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(1.05), Inches(0.32))
        fill(box, color); no_line(box)
        tx(slide, x + 0.04, y + 0.07, 0.95, 0.12, label, 8.5, WHITE, True, PP_ALIGN.CENTER)
    card(slide, 8.22, 2.18, 4.25, 4.55, CARD2)
    bullets(slide, 8.58, 2.85, 3.45, 2.5, ["趋势：现在大概往哪走？", "位置：是在高位、低位还是中间？", "量能：这次变化有没有人参与？"], 17)
    tx(slide, 8.52, 5.75, 3.55, 0.42, "形态不是答案，只是提醒你去问问题。", 15, TEXT, True, PP_ALIGN.CENTER)


def plan(page):
    slide = base("买之前，只问三个问题", page)
    big_line(slide, "不会回答这三个问题，就先别急着买。")
    steps = [("为什么买？", "因为规则/趋势/价格位置"), ("错了怎么办？", "跌到哪里我认错"), ("赚了怎么卖？", "涨到哪里分批走")]
    for i, (a, b) in enumerate(steps):
        x = 1.05 + i * 4.05
        box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.75), Inches(3.35), Inches(1.45))
        box.adjustments[0] = 0.08
        fill(box, [BLUE, RED, GREEN][i]); no_line(box)
        tx(slide, x + 0.18, 3.05, 2.95, 0.3, a, 22, WHITE, True, PP_ALIGN.CENTER)
        tx(slide, x + 0.18, 3.55, 2.95, 0.22, b, 12, WHITE, False, PP_ALIGN.CENTER)
    tx(slide, 1.25, 5.68, 10.8, 0.35, "新手最大的进步：从“我感觉会涨”，变成“我有计划，也知道错了怎么办”。", 19, WHITE, True, PP_ALIGN.CENTER)


def recap(page):
    slide = base("复盘：每次交易都留一条记录", page)
    big_line(slide, "复盘不是写作文，是找出自己哪里做错了。")
    rows = [("买入理由", "我为什么买？"), ("卖出理由", "我为什么卖？"), ("赚亏结果", "赚/亏多少钱？"), ("错误记录", "是追高、乱加仓，还是没止损？")]
    for i, (a, b) in enumerate(rows):
        y = 2.45 + i * 0.86
        card(slide, 1.35, y, 10.65, 0.58, WHITE if i % 2 == 0 else RGBColor(240, 247, 251))
        tx(slide, 1.65, y + 0.14, 2.0, 0.2, a, 14, [BLUE, GREEN, GOLD, RED][i], True)
        tx(slide, 3.85, y + 0.14, 7.6, 0.2, b, 14, TEXT)


def summary(page):
    slide = base("最后只记住这几句话", page)
    big_line(slide, "先会听懂，再去学习更复杂的技术。")
    card(slide, 1.15, 2.30, 11.05, 3.85, WHITE)
    bullets(slide, 1.72, 2.82, 10.0, 2.75, [
        "买入 = 花钱换股票；卖出 = 把股票换回钱。",
        "港股/美股当天买了可以当天卖，交收是另一件事。",
        "K线看价格过程，均线看趋势，成交量看热闹程度。",
    ], 20)
    tx(slide, 1.72, 5.72, 10.0, 0.28, "任何股票案例都不是推荐买入。", 18, RED, True, PP_ALIGN.CENTER)


def build():
    page = 1
    cover(); page += 1
    simple("股票是什么", "股票就是公司的一小份所有权。", ["买股票 = 买这家公司的一小份。", "股价每天会变。", "变贵可能赚钱，变便宜可能亏钱。"], page, BLUE); page += 1
    simple("什么叫买入", "买入 = 你花钱，把股票买到自己账户里。", ["你付出现金。", "账户里多了股票。", "买入价就是你买到的价格。"], page, BLUE); page += 1
    simple("什么叫卖出", "卖出 = 你把股票卖掉，换回现金。", ["你交出股票。", "账户里多了现金。", "卖出价就是你卖掉的价格。"], page, GREEN); page += 1
    math_profit(page); page += 1
    loss_math(page); page += 1
    account(page); page += 1
    order(page); page += 1
    simple("限价单和市价单", "限价单管价格，市价单求速度。", ["限价单：不到这个价就不成交。", "市价单：尽快成交，但价格可能不理想。", "新手优先学会用限价单。"], page, GOLD); page += 1
    deal(page); page += 1
    order_book(page); page += 1
    priority(page); page += 1
    cancel_order(page); page += 1
    price_cage(page); page += 1
    market_rules(page); page += 1
    trading_phases(page); page += 1
    price_limit(page); page += 1
    breakers(page); page += 1
    halt_st(page); page += 1
    dividend(page); page += 1
    chart_slide("真实行情图长什么样", "行情图就是把价格变化画出来。", "600519_SH.png", ["上面：价格走势。", "中间彩线：均线。", "下面：成交量。"], page); page += 1
    kline_single(page); page += 1
    kline_gallery(page); page += 1
    kline_location(page); page += 1
    kline_context(page); page += 1
    chart_slide("均线怎么看", "均线 = 过去一段时间的平均价格。", "NVDA.png", ["短均线：反应快。", "长均线：看大方向。", "均线是辅助，不是保证。"], page); page += 1
    chart_slide("成交量怎么看", "成交量 = 这段时间买卖有多热闹。", "300750_SZ.png", ["量大：参与的人多。", "量小：交易比较冷。", "涨跌都要配合成交量看。"], page); page += 1
    simple("为什么会涨跌", "买的人更着急，价格容易涨；卖的人更着急，价格容易跌。", ["消息会影响情绪。", "业绩会影响预期。", "资金进出会影响价格。"], page, RED); page += 1
    plan(page); page += 1
    simple("仓位和止损", "先决定最多亏多少，再决定买多少。", ["不要一上来满仓。", "亏损后不要随意加仓。", "跌到计划位置要认错。"], page, GOLD); page += 1
    recap(page); page += 1
    summary(page); page += 1
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
